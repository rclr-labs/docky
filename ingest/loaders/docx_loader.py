"""
DOCX Loader: extrai texto de documentos Word (.docx).

Aplicado a atas de reuniao, notas geradas por IA (ex.: "Anotacoes do Gemini"
exportadas do Google Meet) e quaisquer documentos Word do acervo.
Preserva a estrutura de paragrafos e extrai tabelas como linhas "celula | celula".
"""
import logging
from pathlib import Path

from ingest.loader import BaseLoader, Document, DocumentMetadata


logger = logging.getLogger(__name__)

# Chunks maiores nao precisam ser feitos aqui: o DocumentIngestor usa este texto
# e o SearchService aplica chunking (512 chars) quando nao ha chunks prontos.


class DOCXLoader(BaseLoader):
    """Loader para documentos Word (.docx)"""

    async def validate(self, file_path: str) -> bool:
        path = Path(file_path)
        return path.suffix.lower() == ".docx" and path.is_file()

    async def load(self, file_path: str) -> Document:
        logger.info(f"Carregando DOCX: {file_path}")
        from docx import Document as DocxDocument  # import tardio

        path = Path(file_path)
        docx = DocxDocument(str(path))

        parts: list[str] = []
        for para in docx.paragraphs:
            text = para.text.strip()
            if text:
                # Preservar titulos com prefixo markdown-like ajuda o chunking
                # a nao quebrar no meio de secoes
                style = (para.style.name or "").lower()
                if style.startswith("heading"):
                    level = "".join(c for c in style if c.isdigit()) or "2"
                    parts.append("#" * min(int(level), 4) + " " + text)
                else:
                    parts.append(text)

        for table in docx.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        content = "\n\n".join(parts)
        logger.info(f"  {len(parts)} blocos de texto extraidos ({len(content)} chars)")

        metadata = DocumentMetadata(
            file_path=str(path),
            file_name=path.name,
            file_size=path.stat().st_size,
            file_type=".docx",
        )
        return Document(content=content, metadata=metadata)

